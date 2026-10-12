import io
import csv
import logging
from datetime import timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.utils import timezone
from django.core.signing import TimestampSigner, BadSignature, SignatureExpired
from django.core.mail import send_mail
from django.conf import settings
from django.db import transaction

from apps.accounts.models import User, Role

logger = logging.getLogger(__name__)

SETUP_TOKEN_SALT = 'tenant-user-setup'
TOKEN_MAX_AGE_SECONDS = 7 * 86400  # Exactly 7 days (168 hours)


class UserImportService:
    """
    Handles bulk user onboarding from 3-column CSV/Excel files (Name, Mobile Number, Email).
    Generates cryptographically signed, one-time 7-day password setup links.
    """

    @classmethod
    def generate_sample_template(cls, file_format: str = 'csv') -> tuple[io.BytesIO, str, str]:
        """
        Generates a pre-formatted sample template with columns:
        [Name, Mobile Number, Email]
        """
        sample_rows = [
            ["Dr. Ramesh Sharma", "9876543210", "ramesh.sharma@institution.edu"],
            ["Prof. Sunita Patel", "9123456780", "sunita.patel@institution.edu"],
            ["Dr. Amit Verma", "9988776655", "amit.verma@institution.edu"],
        ]
        headers = ["Name", "Mobile Number", "Email"]

        if file_format.lower() == 'xlsx':
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Users Import Template"

            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            center_align = Alignment(horizontal="center", vertical="center")
            left_align = Alignment(horizontal="left", vertical="center")

            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align

            for r_idx, row in enumerate(sample_rows, start=2):
                ws.append(row)
                ws.cell(row=r_idx, column=1).alignment = left_align
                ws.cell(row=r_idx, column=2).alignment = center_align
                ws.cell(row=r_idx, column=3).alignment = left_align

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return buf, "sample_users_template.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        else:  # CSV default
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(headers)
            for row in sample_rows:
                writer.writerow(row)
            bytes_buf = io.BytesIO(buf.getvalue().encode('utf-8'))
            bytes_buf.seek(0)
            return bytes_buf, "sample_users_template.csv", "text/csv"

    @classmethod
    def generate_setup_token(cls, user_id) -> tuple[str, timezone.datetime]:
        """Generates a secure timestamped token valid for 7 days."""
        signer = TimestampSigner(salt=SETUP_TOKEN_SALT)
        token = signer.sign(str(user_id))
        expires_at = timezone.now() + timedelta(seconds=TOKEN_MAX_AGE_SECONDS)
        return token, expires_at

    @classmethod
    def verify_setup_token(cls, user_id_str: str, token: str) -> bool:
        """Verifies if the token signature matches and is within the 7-day window."""
        signer = TimestampSigner(salt=SETUP_TOKEN_SALT)
        try:
            val = signer.unsign(token, max_age=TOKEN_MAX_AGE_SECONDS)
            return val == str(user_id_str)
        except (BadSignature, SignatureExpired):
            return False

    @classmethod
    def parse_file(cls, uploaded_file) -> list[dict]:
        """
        Parses CSV or Excel file content, normalizing header columns:
        Expects: 'Name', 'Mobile Number' (or 'mob no'), 'Email'
        """
        filename = getattr(uploaded_file, 'name', '').lower()
        rows = []

        if filename.endswith('.xlsx') or filename.endswith('.xls'):
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            ws = wb.active
            iter_rows = list(ws.iter_rows(values_only=True))
            if not iter_rows:
                return []
            raw_headers = [str(h or '').strip().lower() for h in iter_rows[0]]
            header_map = cls._map_headers(raw_headers)

            for raw_row in iter_rows[1:]:
                if not any(raw_row):
                    continue
                row_dict = {}
                for key, col_idx in header_map.items():
                    val = raw_row[col_idx] if col_idx < len(raw_row) else ""
                    row_dict[key] = str(val or '').strip()
                if row_dict.get('email'):
                    rows.append(row_dict)
        else:
            # CSV parsing
            content = uploaded_file.read()
            if isinstance(content, bytes):
                try:
                    text = content.decode('utf-8-sig')
                except UnicodeDecodeError:
                    text = content.decode('latin-1')
            else:
                text = str(content)

            reader = csv.reader(io.StringIO(text))
            rows_list = list(reader)
            if not rows_list:
                return []
            raw_headers = [str(h or '').strip().lower() for h in rows_list[0]]
            header_map = cls._map_headers(raw_headers)

            for raw_row in rows_list[1:]:
                if not any(raw_row):
                    continue
                row_dict = {}
                for key, col_idx in header_map.items():
                    val = raw_row[col_idx] if col_idx < len(raw_row) else ""
                    row_dict[key] = str(val or '').strip()
                if row_dict.get('email'):
                    rows.append(row_dict)

        return rows

    @classmethod
    def _map_headers(cls, headers: list[str]) -> dict:
        """Finds column indices for name, mobile, email with generous aliases."""
        header_map = {}
        for idx, h in enumerate(headers):
            clean = h.replace('_', ' ').replace('-', ' ').strip().lower()
            if clean in ('name', 'full name', 'faculty name', 'user name', 'teacher name'):
                header_map['name'] = idx
            elif clean in ('mobile', 'mobile number', 'mob no', 'mob number', 'phone', 'phone number', 'contact'):
                header_map['mobile'] = idx
            elif clean in ('email', 'email id', 'email address', 'official email'):
                header_map['email'] = idx

        # Fallbacks by position if headers are not exact
        if 'email' not in header_map and len(headers) >= 3:
            header_map = {'name': 0, 'mobile': 1, 'email': 2}

        return header_map

    @classmethod
    def bulk_import_users(
        cls,
        uploaded_file,
        tenant,
        creator_user=None,
        base_url: str = "http://localhost:8000"
    ) -> dict:
        """
        Parses rows, validates data, creates user accounts with unusable passwords,
        and generates 7-day one-time setup tokens sent via email.
        """
        parsed_rows = cls.parse_file(uploaded_file)
        if not parsed_rows:
            return {
                "success": False,
                "message": "No valid data rows found in the uploaded file.",
                "total": 0,
                "created": 0,
                "skipped": 0,
                "errors": ["File was empty or headers were unrecognized."]
            }

        created_users = []
        skipped_users = []
        errors = []
        setup_links = []

        default_faculty_role = Role.objects.filter(
            tenant=tenant, code='faculty'
        ).first() or Role.objects.filter(
            tenant__isnull=True, code='faculty'
        ).first()

        for line_num, row in enumerate(parsed_rows, start=2):
            raw_email = row.get('email', '').strip().lower()
            raw_name = row.get('name', '').strip()
            raw_phone = row.get('mobile', '').strip()

            if not raw_email or '@' not in raw_email:
                errors.append(f"Row {line_num}: Invalid email address '{raw_email}'.")
                continue

            if User.objects.filter(email__iexact=raw_email).exists():
                skipped_users.append(f"{raw_email} (already exists)")
                continue

            # Split name into first and last
            parts = raw_name.split() if raw_name else []
            first_name = parts[0] if parts else ""
            last_name = " ".join(parts[1:]) if len(parts) > 1 else ""

            # Derive username
            username_base = raw_email.split('@')[0]
            username = username_base
            counter = 1
            while User.objects.filter(username__iexact=username).exists():
                username = f"{username_base}{counter}"
                counter += 1

            try:
                with transaction.atomic():
                    user = User(
                        tenant=tenant,
                        email=raw_email,
                        username=username,
                        first_name=first_name,
                        last_name=last_name,
                        phone=raw_phone,
                        role=default_faculty_role,
                        is_staff=False,
                        is_active=True,
                        is_approved=True,
                        is_superuser=False,
                    )
                    user.set_unusable_password()
                    user.save()

                    token, expires_at = cls.generate_setup_token(user.id)
                    user.password_setup_token = token
                    user.password_setup_expires_at = expires_at
                    user.save(update_fields=['password_setup_token', 'password_setup_expires_at'])

                    setup_url = f"{base_url}/auth/setup-password?uid={user.id}&token={token}"
                    setup_links.append({"email": user.email, "setup_url": setup_url})

                    # Send Invitation Email
                    cls.send_invitation_email(user, setup_url, tenant)
                    created_users.append(user)

            except Exception as e:
                logger.error(f"Error creating user {raw_email}: {e}")
                errors.append(f"Row {line_num}: Failed to create user '{raw_email}': {str(e)}")

        return {
            "success": True,
            "message": f"Successfully imported {len(created_users)} users.",
            "total": len(parsed_rows),
            "created": len(created_users),
            "skipped": len(skipped_users),
            "skipped_details": skipped_users,
            "errors": errors,
            "setup_links": setup_links[:10],  # Return first 10 for administrative confirmation
        }

    @classmethod
    def send_invitation_email(cls, user, setup_url: str, tenant):
        """Sends the one-time activation & password setup link (valid for 1 week)."""
        tenant_name = tenant.name if tenant else "Poornima University Portal"
        subject = f"Welcome to {tenant_name} — Set Up Your Account Password"
        body = (
            f"Dear {user.get_full_name() or user.username},\n\n"
            f"An account has been created for you on the {tenant_name} Achievements Portal.\n\n"
            f"Please click the link below to set your password and activate your account:\n"
            f"{setup_url}\n\n"
            f"This one-time link is valid for 7 days (1 week). If you do not set your password within this time, "
            f"please contact your institution's administrator to request a new invitation.\n\n"
            f"Best regards,\n"
            f"{tenant_name} Administration"
        )
        try:
            from_email = getattr(settings, 'DEFAULT_FROM_EMAIL', 'no-reply@poornima.org')
            send_mail(subject, body, from_email, [user.email], fail_silently=True)
        except Exception as e:
            logger.warning(f"Could not dispatch email to {user.email}: {e}")

    @classmethod
    def complete_password_setup(cls, user_id_str: str, token: str, new_password: str) -> tuple[bool, str, User | None]:
        """
        Validates the 7-day token and establishes the user's password.
        """
        user = User.objects.filter(id=user_id_str).first()
        if not user:
            return False, "User account not found.", None

        if not user.password_setup_token or user.password_setup_token != token:
            return False, "Invalid or already-consumed password setup token.", None

        if not user.password_setup_expires_at or user.password_setup_expires_at < timezone.now():
            return False, "This one-time setup link has expired (links are valid for 7 days). Please contact your administrator.", None

        if not cls.verify_setup_token(user.id, token):
            return False, "Token signature verification failed or token has expired.", None

        if not new_password or len(new_password) < 8:
            return False, "Password must be at least 8 characters long.", None

        # Set password and invalidate token
        with transaction.atomic():
            user.set_password(new_password)
            user.password_setup_token = None
            user.password_setup_expires_at = None
            user.is_active = True
            user.is_approved = True
            user.save(update_fields=['password', 'password_setup_token', 'password_setup_expires_at', 'is_active', 'is_approved'])

        return True, "Password created successfully. You may now log in.", user
