import io
from urllib.parse import urljoin
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.conf import settings
from django.db.models import Q
from apps.activities.models import Activity, ActivityStatus, ActivityType


class AccreditationExportService:
    """
    Generates standardized, multi-tab Excel workbooks formatted strictly
    for official NAAC AQAR Criteria 3 and NIRF institutional evaluations.
    """

    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Deep navy
    DATA_FONT = Font(name="Calibri", size=10)
    LINK_FONT = Font(name="Calibri", size=10, color="0000FF", underline="single")
    CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
    THIN_BORDER = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    @classmethod
    def resolve_proof_url(cls, document) -> str:
        """Resolves direct download / preview URL for an attachment."""
        if not document or not document.file:
            return ""
        try:
            url = document.direct_url
            if url.startswith("http://") or url.startswith("https://"):
                return url
            public_base = getattr(settings, 'PUBLIC_MEDIA_URL', 'http://localhost:8000/media/')
            return urljoin(public_base, document.file.name)
        except Exception:
            return ""

    @classmethod
    def get_meta(cls, m: dict, *candidates, default=""):
        """Extracts values from polymorphic metadata dictionary with alias fallback and key normalization."""
        if not m or not isinstance(m, dict):
            return default
        for cand in candidates:
            if cand in m and m[cand] not in (None, ""):
                return str(m[cand]).strip()
        norm_map = {}
        for k, v in m.items():
            if v not in (None, ""):
                k_clean = str(k).lower().strip().replace('_', '').replace(' ', '')
                norm_map[k_clean] = str(v).strip()
        for cand in candidates:
            cand_clean = str(cand).lower().strip().replace('_', '').replace(' ', '')
            if cand_clean in norm_map:
                return norm_map[cand_clean]
        return default

    @classmethod
    def generate_naac_criteria_3_workbook(
        cls,
        academic_session_id=None,
        department_id=None,
        faculty_id=None,
        statuses=None,
        verified=None,
        has_proof=None,
        activity_types=None,
        date_from=None,
        date_to=None,
        search=None,
        activity_ids=None,
        tenant_id=None,
    ) -> io.BytesIO:
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Base query
        base_qs = Activity.objects.all().select_related(
            'faculty', 'faculty__department', 'faculty__designation', 'academic_session', 'tenant'
        ).prefetch_related('documents', 'activity_authors')

        # 0. Tenant isolation
        if tenant_id:
            base_qs = base_qs.filter(tenant_id=tenant_id)

        # 1. Status filtering
        if statuses in ('all', '*'):
            pass  # Include all statuses
        elif isinstance(statuses, (list, tuple, set)) and statuses:
            base_qs = base_qs.filter(status__in=[s.lower() for s in statuses])
        elif isinstance(statuses, str) and statuses.strip():
            st_list = [s.strip().lower() for s in statuses.split(',') if s.strip()]
            if 'all' not in st_list and '*' not in st_list:
                base_qs = base_qs.filter(status__in=st_list)
        else:
            base_qs = base_qs.filter(status=ActivityStatus.APPROVED)

        # 2. Verification filtering
        if verified is not None and str(verified).lower() not in ('', 'all', '*'):
            v_str = str(verified).lower()
            if v_str in ('true', '1', 'yes', 'verified'):
                base_qs = base_qs.filter(
                    Q(verified_by__isnull=False) |
                    Q(status=ActivityStatus.APPROVED) |
                    Q(documents__is_verified=True)
                ).distinct()
            elif v_str in ('false', '0', 'no', 'unverified'):
                base_qs = base_qs.filter(
                    verified_by__isnull=True
                ).exclude(
                    status=ActivityStatus.APPROVED
                ).exclude(
                    documents__is_verified=True
                ).distinct()

        # 3. Proof attachment filtering
        if has_proof is not None and str(has_proof).lower() not in ('', 'all', '*'):
            hp_str = str(has_proof).lower()
            if hp_str in ('true', '1', 'yes', 'with_proof'):
                base_qs = base_qs.filter(documents__isnull=False).distinct()
            elif hp_str in ('false', '0', 'no', 'without_proof'):
                base_qs = base_qs.filter(documents__isnull=True)

        # 4. Academic session filtering
        if academic_session_id and str(academic_session_id).lower() not in ('', 'all', '*'):
            base_qs = base_qs.filter(
                Q(academic_session_id=academic_session_id) |
                Q(academic_session__name__iexact=str(academic_session_id))
            )

        # 5. Department filtering
        if department_id and str(department_id).lower() not in ('', 'all', '*'):
            base_qs = base_qs.filter(
                Q(faculty__department_id=department_id) |
                Q(faculty__department__code__iexact=str(department_id)) |
                Q(faculty__department__name__iexact=str(department_id))
            )

        # 6. Faculty filtering
        if faculty_id and str(faculty_id).lower() not in ('', 'all', '*'):
            base_qs = base_qs.filter(
                Q(faculty_id=faculty_id) |
                Q(faculty__employee_id__iexact=str(faculty_id)) |
                Q(faculty__email__iexact=str(faculty_id))
            )

        # 7. Date range filtering
        if date_from:
            try:
                base_qs = base_qs.filter(created_at__date__gte=str(date_from)[:10])
            except Exception:
                pass
        if date_to:
            try:
                base_qs = base_qs.filter(created_at__date__lte=str(date_to)[:10])
            except Exception:
                pass

        # 8. Search keyword filtering
        if search:
            s = str(search).strip()
            base_qs = base_qs.filter(
                Q(title__icontains=s) |
                Q(faculty__first_name__icontains=s) |
                Q(faculty__last_name__icontains=s) |
                Q(faculty__employee_id__icontains=s)
            )

        # 9. Specific activity IDs filtering
        if activity_ids:
            if isinstance(activity_ids, str):
                id_list = [i.strip() for i in activity_ids.split(',') if i.strip()]
            else:
                id_list = list(activity_ids)
            base_qs = base_qs.filter(id__in=id_list)

        # 10. Determine which sheets to include
        types_set = None
        if activity_types and activity_types not in ('all', '*'):
            if isinstance(activity_types, str):
                types_set = {t.strip().lower() for t in activity_types.split(',') if t.strip()}
            else:
                types_set = {t.lower() for t in activity_types}
            if 'all' in types_set or '*' in types_set:
                types_set = None

        include_journals = types_set is None or ActivityType.JOURNAL_PUBLICATION in types_set
        include_conf = types_set is None or ActivityType.CONFERENCE_PUBLICATION in types_set
        include_patents = types_set is None or ActivityType.PATENT in types_set
        include_part = types_set is None or ActivityType.FACULTY_PARTICIPATION in types_set

        # ── Tab 1: NAAC 3.4.3 (Journals) ──────────────────────────────────────
        if include_journals:
            cls._write_journals_sheet(wb, base_qs.filter(activity_type=ActivityType.JOURNAL_PUBLICATION).order_by('-created_at'))

        # ── Tab 2: NAAC 3.4.4 (Conferences) ───────────────────────────────────
        if include_conf:
            cls._write_conferences_sheet(wb, base_qs.filter(activity_type=ActivityType.CONFERENCE_PUBLICATION).order_by('-created_at'))

        # ── Tab 3: NAAC 3.4.5 (Patents) ───────────────────────────────────────
        if include_patents:
            cls._write_patents_sheet(wb, base_qs.filter(activity_type=ActivityType.PATENT).order_by('-created_at'))

        # ── Tab 4: NAAC 3.4.6 (Faculty Participation / FDPs) ──────────────────
        if include_part:
            cls._write_participation_sheet(wb, base_qs.filter(activity_type=ActivityType.FACULTY_PARTICIPATION).order_by('-created_at'))

        if len(wb.sheetnames) == 0:
            ws_empty = wb.create_sheet(title="NAAC Criteria 3")
            ws_empty.cell(row=1, column=1, value="No activity records found matching requested criteria.")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def _write_journals_sheet(cls, wb, journal_qs):
        ws_journals = wb.create_sheet(title="NAAC 3.4.3 (Journals)")
        journal_headers = [
            "S. No.", "Name of the Teacher", "Department", "Employee ID",
            "Current Designation", "Title of the Paper", "Name of the Journal",
            "Year of Publication", "ISSN", "Indexing (Scopus/WoS/UGC/Other)",
            "DOI / Paper URL", "Link to Proof Document"
        ]
        cls._write_sheet_headers(ws_journals, journal_headers)

        row_idx = 2
        for idx, act in enumerate(journal_qs, 1):
            m = act.metadata or {}
            proof_doc = act.documents.first()
            proof_url = cls.resolve_proof_url(proof_doc)
            doi_or_url = cls.get_meta(m, 'doi', 'paper_url')

            teacher_name = act.faculty.display_name or cls.get_meta(m, 'name_of_faculty')
            dept_name = (act.faculty.department.name if act.faculty.department else "") or cls.get_meta(m, 'deparment', 'department')
            emp_id = (act.faculty.employee_id or "") or cls.get_meta(m, 'emp_id', 'employee_id')
            desig_name = (act.faculty.designation.name if act.faculty.designation else "") or cls.get_meta(m, 'current_designation', 'designation')
            journal_name = cls.get_meta(m, 'name_of_journal', 'journal_name', 'journal_name_book_name_procedings_name_(')
            raw_pub_year = cls.get_meta(m, 'publication_year', 'date_of_publication', 'publication_date')
            pub_year = str(raw_pub_year)[:4] if raw_pub_year else ""
            issn = cls.get_meta(m, 'isbn_issn_(if_applicable)', 'issn', 'isbn')
            indexing = cls.get_meta(m, 'indexing(if_any)_eg_scopus,_sci,_ugc_car', 'indexing', default='Peer Reviewed')

            ws_journals.cell(row=row_idx, column=1, value=idx).alignment = cls.CENTER_ALIGN
            ws_journals.cell(row=row_idx, column=2, value=teacher_name).alignment = cls.LEFT_ALIGN
            ws_journals.cell(row=row_idx, column=3, value=dept_name).alignment = cls.LEFT_ALIGN
            ws_journals.cell(row=row_idx, column=4, value=emp_id).alignment = cls.CENTER_ALIGN
            ws_journals.cell(row=row_idx, column=5, value=desig_name).alignment = cls.LEFT_ALIGN
            ws_journals.cell(row=row_idx, column=6, value=act.title).alignment = cls.LEFT_ALIGN
            ws_journals.cell(row=row_idx, column=7, value=journal_name).alignment = cls.LEFT_ALIGN
            ws_journals.cell(row=row_idx, column=8, value=pub_year).alignment = cls.CENTER_ALIGN
            ws_journals.cell(row=row_idx, column=9, value=issn).alignment = cls.CENTER_ALIGN
            ws_journals.cell(row=row_idx, column=10, value=indexing).alignment = cls.CENTER_ALIGN

            # DOI
            cell_doi = ws_journals.cell(row=row_idx, column=11, value=doi_or_url)
            cell_doi.alignment = cls.LEFT_ALIGN

            # Proof URL
            cell_proof = ws_journals.cell(row=row_idx, column=12)
            if proof_url:
                cell_proof.value = "View Proof"
                cell_proof.hyperlink = proof_url
                cell_proof.font = cls.LINK_FONT
            else:
                cell_proof.value = "No Attachment"
                cell_proof.font = cls.DATA_FONT
            cell_proof.alignment = cls.CENTER_ALIGN

            cls._apply_row_styling(ws_journals, row_idx, len(journal_headers))
            row_idx += 1

        cls._auto_fit_columns(ws_journals)

    @classmethod
    def _write_conferences_sheet(cls, wb, conf_qs):
        ws_conf = wb.create_sheet(title="NAAC 3.4.4 (Conferences)")
        conf_headers = [
            "S. No.", "Name of the Teacher", "Department", "Employee ID",
            "Title of the Paper", "Title of the Proceedings", "Name of Conference",
            "Year of Publication", "ISBN", "Affiliating Institute",
            "Publisher", "Link to Proof Document"
        ]
        cls._write_sheet_headers(ws_conf, conf_headers)

        row_idx = 2
        for idx, act in enumerate(conf_qs, 1):
            m = act.metadata or {}
            proof_doc = act.documents.first()
            proof_url = cls.resolve_proof_url(proof_doc)

            teacher_name = act.faculty.display_name or cls.get_meta(m, 'name_of_faculty')
            dept_name = (act.faculty.department.name if act.faculty.department else "") or cls.get_meta(m, 'deparment', 'department')
            emp_id = (act.faculty.employee_id or "") or cls.get_meta(m, 'emp_id', 'employee_id')
            proceedings = cls.get_meta(m, 'journal_name_book_name_procedings_name_(', 'proceedings_title', 'title_of_the_proceedings')
            conf_name = cls.get_meta(m, 'conference_name', 'name_of_conference', 'types_of_conference_(national_conference')
            raw_conf_year = cls.get_meta(m, 'conference_year', 'date_of_conference', 'conference_date')
            conf_year = str(raw_conf_year)[:4] if raw_conf_year else ""
            isbn = cls.get_meta(m, 'isbn_issn_(if_applicable)', 'isbn', 'issn')
            affiliating = cls.get_meta(m, 'affiliating_institute_at_the_time_of_pub', 'affiliating_institute', default='Poornima College of Engineering')
            publisher = cls.get_meta(m, 'name_of_the_publisher', 'publisher')

            ws_conf.cell(row=row_idx, column=1, value=idx).alignment = cls.CENTER_ALIGN
            ws_conf.cell(row=row_idx, column=2, value=teacher_name).alignment = cls.LEFT_ALIGN
            ws_conf.cell(row=row_idx, column=3, value=dept_name).alignment = cls.LEFT_ALIGN
            ws_conf.cell(row=row_idx, column=4, value=emp_id).alignment = cls.CENTER_ALIGN
            ws_conf.cell(row=row_idx, column=5, value=act.title).alignment = cls.LEFT_ALIGN
            ws_conf.cell(row=row_idx, column=6, value=proceedings).alignment = cls.LEFT_ALIGN
            ws_conf.cell(row=row_idx, column=7, value=conf_name).alignment = cls.LEFT_ALIGN
            ws_conf.cell(row=row_idx, column=8, value=conf_year).alignment = cls.CENTER_ALIGN
            ws_conf.cell(row=row_idx, column=9, value=isbn).alignment = cls.CENTER_ALIGN
            ws_conf.cell(row=row_idx, column=10, value=affiliating).alignment = cls.LEFT_ALIGN
            ws_conf.cell(row=row_idx, column=11, value=publisher).alignment = cls.LEFT_ALIGN

            cell_proof = ws_conf.cell(row=row_idx, column=12)
            if proof_url:
                cell_proof.value = "View Proof"
                cell_proof.hyperlink = proof_url
                cell_proof.font = cls.LINK_FONT
            else:
                cell_proof.value = "No Attachment"
                cell_proof.font = cls.DATA_FONT
            cell_proof.alignment = cls.CENTER_ALIGN

            cls._apply_row_styling(ws_conf, row_idx, len(conf_headers))
            row_idx += 1

        cls._auto_fit_columns(ws_conf)

    @classmethod
    def _write_patents_sheet(cls, wb, patent_qs):
        ws_patents = wb.create_sheet(title="NAAC 3.4.5 (Patents)")
        patent_headers = [
            "S. No.", "Name of the Teacher", "Department", "Employee ID",
            "Title of Patent", "Patent / Application Number", "Filing Date",
            "Grant Date / Publication Date", "Status (Filed/Published/Granted)",
            "Country", "Link to Proof Document"
        ]
        cls._write_sheet_headers(ws_patents, patent_headers)

        row_idx = 2
        for idx, act in enumerate(patent_qs, 1):
            m = act.metadata or {}
            proof_doc = act.documents.first()
            proof_url = cls.resolve_proof_url(proof_doc)

            teacher_name = act.faculty.display_name or cls.get_meta(m, 'name_of_faculty')
            dept_name = (act.faculty.department.name if act.faculty.department else "") or cls.get_meta(m, 'deparment', 'department')
            emp_id = (act.faculty.employee_id or "") or cls.get_meta(m, 'emp_id', 'employee_id')
            patent_no = cls.get_meta(m, 'application_number', 'patent_or_application_number', 'patent_number')
            filing_date = str(cls.get_meta(m, 'filing_date', 'date_of_filing'))[:10]
            pub_date = str(cls.get_meta(m, 'date_of_publication', 'publication_date', 'grant_date'))[:10]
            patent_status = cls.get_meta(m, 'patent_status', default='Published')
            country = cls.get_meta(m, 'country', default='India')

            ws_patents.cell(row=row_idx, column=1, value=idx).alignment = cls.CENTER_ALIGN
            ws_patents.cell(row=row_idx, column=2, value=teacher_name).alignment = cls.LEFT_ALIGN
            ws_patents.cell(row=row_idx, column=3, value=dept_name).alignment = cls.LEFT_ALIGN
            ws_patents.cell(row=row_idx, column=4, value=emp_id).alignment = cls.CENTER_ALIGN
            ws_patents.cell(row=row_idx, column=5, value=act.title).alignment = cls.LEFT_ALIGN
            ws_patents.cell(row=row_idx, column=6, value=patent_no).alignment = cls.CENTER_ALIGN
            ws_patents.cell(row=row_idx, column=7, value=filing_date).alignment = cls.CENTER_ALIGN
            ws_patents.cell(row=row_idx, column=8, value=pub_date).alignment = cls.CENTER_ALIGN
            ws_patents.cell(row=row_idx, column=9, value=patent_status).alignment = cls.CENTER_ALIGN
            ws_patents.cell(row=row_idx, column=10, value=country).alignment = cls.CENTER_ALIGN

            cell_proof = ws_patents.cell(row=row_idx, column=11)
            if proof_url:
                cell_proof.value = "View Proof"
                cell_proof.hyperlink = proof_url
                cell_proof.font = cls.LINK_FONT
            else:
                cell_proof.value = "No Attachment"
                cell_proof.font = cls.DATA_FONT
            cell_proof.alignment = cls.CENTER_ALIGN

            cls._apply_row_styling(ws_patents, row_idx, len(patent_headers))
            row_idx += 1

        cls._auto_fit_columns(ws_patents)

    @classmethod
    def _write_participation_sheet(cls, wb, fdp_qs):
        ws_fdp = wb.create_sheet(title="NAAC 3.4.6 (Participation)")
        fdp_headers = [
            "Sr. No.", "Name of Faculty", "Department", "Employee ID",
            "Type of Program (FDP/Workshop/STTP)", "Title of the Program",
            "Duration (Days)", "Start Date", "End Date", "Organized By",
            "Sponsored By", "Venue", "Link to Proof Document"
        ]
        cls._write_sheet_headers(ws_fdp, fdp_headers)

        row_idx = 2
        for idx, act in enumerate(fdp_qs, 1):
            m = act.metadata or {}
            proof_doc = act.documents.first()
            proof_url = cls.resolve_proof_url(proof_doc)

            teacher_name = act.faculty.display_name or cls.get_meta(m, 'name_of_faculty')
            dept_name = (act.faculty.department.name if act.faculty.department else "") or cls.get_meta(m, 'deparment', 'department')
            emp_id = (act.faculty.employee_id or "") or cls.get_meta(m, 'emp_id', 'employee_id')
            prog_type = cls.get_meta(m, 'type_of_program_(fdp__workshop_conferenc', 'program_type', 'participation_type', default='FDP')
            prog_title = act.title or cls.get_meta(m, 'title_of_the_program', 'program_title')
            duration = cls.get_meta(m, 'duration_of_the_program', 'duration_days', 'duration')
            start_date = str(cls.get_meta(m, 'start_date'))[:10]
            end_date = str(cls.get_meta(m, 'end_date'))[:10]
            org_by = cls.get_meta(m, 'organized_by')
            spon_by = cls.get_meta(m, 'sponsored_by')
            venue = cls.get_meta(m, 'venue_of_the_program', 'venue')

            ws_fdp.cell(row=row_idx, column=1, value=idx).alignment = cls.CENTER_ALIGN
            ws_fdp.cell(row=row_idx, column=2, value=teacher_name).alignment = cls.LEFT_ALIGN
            ws_fdp.cell(row=row_idx, column=3, value=dept_name).alignment = cls.LEFT_ALIGN
            ws_fdp.cell(row=row_idx, column=4, value=emp_id).alignment = cls.CENTER_ALIGN
            ws_fdp.cell(row=row_idx, column=5, value=prog_type).alignment = cls.CENTER_ALIGN
            ws_fdp.cell(row=row_idx, column=6, value=prog_title).alignment = cls.LEFT_ALIGN
            ws_fdp.cell(row=row_idx, column=7, value=duration).alignment = cls.CENTER_ALIGN
            ws_fdp.cell(row=row_idx, column=8, value=start_date).alignment = cls.CENTER_ALIGN
            ws_fdp.cell(row=row_idx, column=9, value=end_date).alignment = cls.CENTER_ALIGN
            ws_fdp.cell(row=row_idx, column=10, value=org_by).alignment = cls.LEFT_ALIGN
            ws_fdp.cell(row=row_idx, column=11, value=spon_by).alignment = cls.LEFT_ALIGN
            ws_fdp.cell(row=row_idx, column=12, value=venue).alignment = cls.LEFT_ALIGN

            cell_proof = ws_fdp.cell(row=row_idx, column=13)
            if proof_url:
                cell_proof.value = "View Proof"
                cell_proof.hyperlink = proof_url
                cell_proof.font = cls.LINK_FONT
            else:
                cell_proof.value = "No Attachment"
                cell_proof.font = cls.DATA_FONT
            cell_proof.alignment = cls.CENTER_ALIGN

            cls._apply_row_styling(ws_fdp, row_idx, len(fdp_headers))
            row_idx += 1

        cls._auto_fit_columns(ws_fdp)

    @classmethod
    def _write_sheet_headers(cls, ws, headers):
        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header_text)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = cls.CENTER_ALIGN
            cell.border = cls.THIN_BORDER
        ws.row_dimensions[1].height = 28

    @classmethod
    def _apply_row_styling(cls, ws, row_idx, col_count):
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = cls.THIN_BORDER
            if not cell.font or cell.font == cls.DATA_FONT:
                cell.font = cls.DATA_FONT
        ws.row_dimensions[row_idx].height = 22

    @classmethod
    def _auto_fit_columns(cls, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
