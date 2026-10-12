import io
import zipfile
from unittest.mock import patch, MagicMock
import openpyxl

from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.accounts.models import User, Department, Designation
from apps.activities.models import AcademicSession, Activity, ActivityAuthor, ActivityStatus, ActivityType
from apps.documents.models import ActivityDocument, DocumentType
from apps.core.services.ingestion_service import BulkIngestionService
from apps.core.services.export_service import AccreditationExportService
from apps.core.services.crossref import CrossRefService


class BulkIngestionServiceTests(TestCase):
    def test_clean_emp_id(self):
        self.assertEqual(BulkIngestionService.clean_emp_id(6450.0), "6450")
        self.assertEqual(BulkIngestionService.clean_emp_id("6450.0"), "6450")
        self.assertEqual(BulkIngestionService.clean_emp_id("EMP100"), "EMP100")
        self.assertEqual(BulkIngestionService.clean_emp_id(None), "")
        self.assertEqual(BulkIngestionService.clean_emp_id(""), "")

    def test_detect_format_from_headers(self):
        journal_headers = ["S. No.", "Name of Faculty", "Title of Publictaion", "Name of Journal", "ISSN"]
        self.assertEqual(BulkIngestionService.detect_format_from_headers(journal_headers), ActivityType.JOURNAL_PUBLICATION)

        conf_headers = ["S. No.", "Name of Faculty", "Title of Publictaion", "Conference Proceedings", "ISBN"]
        self.assertEqual(BulkIngestionService.detect_format_from_headers(conf_headers), ActivityType.CONFERENCE_PUBLICATION)

        patent_headers = ["S. No.", "Name of Faculty", "Title of Patent", "Patent Application Number", "Filing Date"]
        self.assertEqual(BulkIngestionService.detect_format_from_headers(patent_headers), ActivityType.PATENT)

        fdp_headers = ["S. No.", "Name of Teacher", "Title of the Program", "FDP / Workshop Organizers"]
        self.assertEqual(BulkIngestionService.detect_format_from_headers(fdp_headers), ActivityType.FACULTY_PARTICIPATION)

    def test_resolve_or_create_faculty(self):
        user = BulkIngestionService.resolve_or_create_faculty(
            name="Dr. Sunita Rawat",
            emp_id="6450.0",
            dept_raw="Advance Computing",
            desig_raw="Professor"
        )
        self.assertIsNotNone(user)
        self.assertEqual(user.employee_id, "6450")
        self.assertEqual(user.department.name, "Advance Computing")
        self.assertEqual(user.designation.name, "Professor")

        # Resolving same faculty by employee ID should return existing
        user2 = BulkIngestionService.resolve_or_create_faculty(
            name="Sunita Rawat",
            emp_id="6450",
            dept_raw="Advance Computing"
        )
        self.assertEqual(user.id, user2.id)

    def test_process_xlsx_with_stitched_rows_and_placeholders(self):
        """Simulates real Excel with blank placeholder row (S.No=2.0) and stitched multi-line wrap."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Journal Publications"

        # Headers
        headers = [
            "S. No.", "Name of Faculty", "Emp. Id", "Department", "Current Designation",
            "Title of Publictaion", "Name of Journal", "Year of Publication", "ISSN",
            "Authors name in sequence as shown in publication"
        ]
        ws.append(headers)

        # Row 1: Valid Journal Publication
        ws.append([
            1, "Dr. Amit Verma", 5001, "Computer Science", "Associate Professor",
            "Federated Learning in Edge Devices", "IEEE Transactions", 2026, "1234-5678",
            "Dr. Amit Verma, Dr. John Doe"
        ])

        # Row 2: Wrapped continuation line (wrapped author list)
        ws.append([
            "", "", "", "", "",
            "", "", "", "",
            "Dr. Jane Smith, Dr. Alex Brown"
        ])

        # Row 3: Blank placeholder row (just S.No = 2.0, nothing else)
        ws.append([
            2.0, "", "", "", "",
            "", "", "", "", ""
        ])

        # Row 4: Second valid publication
        ws.append([
            3, "Dr. Neha Gupta", 5002, "Computer Science", "Assistant Professor",
            "Explainable AI for Healthcare", "Springer Nature", 2026, "9876-5432",
            "Dr. Neha Gupta"
        ])

        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_buf.seek(0)

        result = BulkIngestionService.process_xlsx(excel_buf)
        self.assertTrue(result["success"])
        self.assertEqual(result["created"], 2)
        self.assertEqual(result["skipped"], 0)

        # Check that Dr. Amit Verma's co-authors were stitched together
        act1 = Activity.objects.filter(faculty__employee_id="5001").first()
        self.assertIsNotNone(act1)
        author_names = list(act1.activity_authors.values_list('author_name', flat=True))
        self.assertIn("Dr. Amit Verma", author_names)
        self.assertIn("Dr. Jane Smith", author_names)

    def test_process_zip_with_proofs(self):
        """Creates an in-memory ZIP with workbook and proof PDF, testing automatic attachment linking."""
        wb = openpyxl.Workbook()
        ws = wb.active
        headers = [
            "S. No.", "Name of Faculty", "Emp. Id", "Department",
            "Title of Publictaion", "Name of Journal", "proof_filename"
        ]
        ws.append(headers)
        ws.append([
            1, "Dr. Vikram Singh", "7001", "Advance Computing",
            "Blockchain Architecture for Smart Cities", "Elsevier Computer Networks", "blockchain_proof.pdf"
        ])

        xlsx_buf = io.BytesIO()
        wb.save(xlsx_buf)
        xlsx_buf.seek(0)

        # Pack into ZIP
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, 'w') as z:
            z.writestr("faculty_data.xlsx", xlsx_buf.getvalue())
            z.writestr("proofs/blockchain_proof.pdf", b"%PDF-1.4 dummy proof data")

        zip_buf.seek(0)
        res = BulkIngestionService.process_zip(zip_buf)

        self.assertTrue(res["success"])
        self.assertEqual(res["activities_created"], 1)
        self.assertEqual(res["documents_linked"], 1)

        # Verify Activity and Document in database
        act = Activity.objects.filter(faculty__employee_id="7001").first()
        self.assertIsNotNone(act)
        self.assertEqual(act.documents.count(), 1)
        doc = act.documents.first()
        self.assertEqual(doc.title, "blockchain_proof.pdf")
        self.assertTrue(doc.is_pdf)


class AccreditationExportServiceTests(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name="Advance Computing", code="AC")
        self.desig = Designation.objects.create(name="Professor", code="PROF")
        self.session = AcademicSession.objects.create(
            name="2025-2026", start_date="2025-07-01", end_date="2026-06-30", is_current=True
        )
        self.faculty = User.objects.create_user(
            username="prof_export",
            employee_id="EXP001",
            department=self.dept,
            designation=self.desig,
            first_name="Sunita",
            last_name="Rawat"
        )

        # Create 1 approved activity of each type with proof attachments
        for act_type, title in [
            (ActivityType.JOURNAL_PUBLICATION, "Deep Learning in Biomedicine"),
            (ActivityType.CONFERENCE_PUBLICATION, "Smart Sensors in IoT"),
            (ActivityType.PATENT, "Automated Solar Tracking Device"),
            (ActivityType.FACULTY_PARTICIPATION, "5-Day Workshop on Quantum Computing")
        ]:
            act = Activity.objects.create(
                faculty=self.faculty,
                academic_session=self.session,
                activity_type=act_type,
                title=title,
                status=ActivityStatus.APPROVED,
                metadata={"doi": "10.1000/182", "publication_year": "2026"}
            )
            dummy_file = SimpleUploadedFile("cert.pdf", b"%PDF-1.4 dummy", content_type="application/pdf")
            ActivityDocument.objects.create(
                activity=act,
                document_type=DocumentType.PUBLISHED_PAPER,
                file=dummy_file
            )

    def test_generate_naac_criteria_3_workbook_structure(self):
        buf = AccreditationExportService.generate_naac_criteria_3_workbook()
        self.assertIsNotNone(buf)
        buf.seek(0)

        wb = openpyxl.load_workbook(buf, data_only=True)
        expected_sheets = [
            "NAAC 3.4.3 (Journals)",
            "NAAC 3.4.4 (Conferences)",
            "NAAC 3.4.5 (Patents)",
            "NAAC 3.4.6 (Participation)"
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Inspect Journals sheet
        ws_j = wb["NAAC 3.4.3 (Journals)"]
        self.assertEqual(ws_j.cell(row=2, column=6).value, "Deep Learning in Biomedicine")
        self.assertEqual(ws_j.cell(row=2, column=12).value, "View Proof")

        # Inspect Conferences sheet
        ws_c = wb["NAAC 3.4.4 (Conferences)"]
        self.assertEqual(ws_c.cell(row=2, column=5).value, "Smart Sensors in IoT")

        # Inspect Patents sheet
        ws_p = wb["NAAC 3.4.5 (Patents)"]
        self.assertEqual(ws_p.cell(row=2, column=5).value, "Automated Solar Tracking Device")

        # Inspect FDPs sheet
        ws_f = wb["NAAC 3.4.6 (Participation)"]
        self.assertEqual(ws_f.cell(row=2, column=6).value, "5-Day Workshop on Quantum Computing")


class CrossRefServiceTests(TestCase):
    def test_clean_doi(self):
        self.assertEqual(CrossRefService.clean_doi("https://doi.org/10.1109/5.771073"), "10.1109/5.771073")
        self.assertEqual(CrossRefService.clean_doi("http://doi.org/10.1109/5.771073"), "10.1109/5.771073")
        self.assertEqual(CrossRefService.clean_doi("doi:10.1109/5.771073"), "10.1109/5.771073")
        self.assertEqual(CrossRefService.clean_doi("10.1109/5.771073"), "10.1109/5.771073")
        self.assertEqual(CrossRefService.clean_doi("  10.1109/5.771073  "), "10.1109/5.771073")

    @patch('requests.get')
    def test_fetch_doi_mock_success(self, mock_get):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {
            "message": {
                "title": ["Artificial Intelligence in Higher Education"],
                "container-title": ["Journal of Educational Technology"],
                "publisher": "Springer",
                "volume": "42",
                "issue": "3",
                "page": "110-125",
                "ISSN": ["2190-8389"],
                "published-print": {"date-parts": [[2026, 3]]},
                "author": [
                    {"given": "Sunita", "family": "Rawat"},
                    {"given": "Amit", "family": "Verma"}
                ]
            }
        }
        mock_get.return_value = mock_resp

        res = CrossRefService.fetch_doi("10.1007/s11528-026-0001-x")
        self.assertTrue(res["success"])
        data = res["data"]
        self.assertEqual(data["title"], "Artificial Intelligence in Higher Education")
        self.assertEqual(data["journal_name"], "Journal of Educational Technology")
        self.assertEqual(data["publication_year"], "2026")
        self.assertEqual(data["authors_text"], "Sunita Rawat, Amit Verma")

    @patch('requests.get')
    def test_fetch_doi_not_found(self, mock_get):
        mock_resp = MagicMock()
        mock_resp.status_code = 404
        mock_get.return_value = mock_resp

        res = CrossRefService.fetch_doi("10.9999/nonexistent.doi")
        self.assertFalse(res["success"])
        self.assertIn("not found", res["error"])


class SystemConfigurationTests(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username="sysadmin",
            email="sysadmin@poornima.org",
            password="password123",
            first_name="Super",
            last_name="Admin"
        )
        self.faculty_user = User.objects.create_user(
            username="regular_fac",
            email="fac@poornima.org",
            password="password123",
            is_staff=True
        )

    def test_singleton_get_solo(self):
        from apps.core.models import SystemConfiguration
        config1 = SystemConfiguration.get_solo()
        config2 = SystemConfiguration.get_solo()
        self.assertEqual(config1.id, config2.id)
        self.assertEqual(config1.institution_name, "Poornima University")

    def test_superuser_can_access_system_configuration_admin(self):
        from django.test import Client
        c = Client()
        c.force_login(self.admin_user)

        # Changelist view redirects directly to change form
        res = c.get("/admin/core/systemconfiguration/")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/change/", res.url)

        # Follow redirect to change page
        res_change = c.get(res.url)
        self.assertEqual(res_change.status_code, 200)
        self.assertContains(res_change, "Accreditation &amp; Institutional Governance")
        self.assertContains(res_change, "File Upload Limits &amp; Guardrails")

    def test_regular_user_cannot_access_system_configuration(self):
        from django.test import Client
        c = Client()
        c.force_login(self.faculty_user)

        res = c.get("/admin/core/systemconfiguration/")
        # Should be 403 Forbidden or redirected away (permission denied)
        self.assertIn(res.status_code, [403, 302])

    def test_runtime_storage_returns_default_storage(self):
        from apps.documents.storage import get_runtime_storage
        from django.core.files.storage import default_storage

        storage = get_runtime_storage()
        self.assertEqual(storage, default_storage)


class CloudflareD1ServiceTests(TestCase):
    def test_d1_config_resolution(self):
        from apps.core.services.d1_service import get_d1_config
        from django.conf import settings

        with patch.object(settings, 'CLOUDFLARE_ACCOUNT_ID', 'acc-12345'), \
             patch.object(settings, 'CLOUDFLARE_D1_DATABASE_ID', 'd1-uuid-67890'), \
             patch.object(settings, 'CLOUDFLARE_D1_API_TOKEN', 'env-cf-token-xyz'):
            d1_cfg = get_d1_config()
            self.assertEqual(d1_cfg["account_id"], "acc-12345")
            self.assertEqual(d1_cfg["database_id"], "d1-uuid-67890")
            self.assertEqual(d1_cfg["api_token"], "env-cf-token-xyz")

    @patch("apps.core.services.d1_service.requests.post")
    def test_execute_d1_query_mocked(self, mock_post):
        from apps.core.services.d1_service import execute_d1_query

        mock_response = MagicMock()
        mock_response.ok = True
        mock_response.json.return_value = {
            "success": True,
            "result": [
                {
                    "results": [{"id": 1, "username": "faculty1"}],
                    "meta": {"duration": 1.2, "rows_read": 1, "rows_written": 0}
                }
            ]
        }
        mock_post.return_value = mock_response

        cfg = {
            "account_id": "acc-12345",
            "database_id": "d1-uuid-67890",
            "api_token": "cf-token-abc",
        }

        res = execute_d1_query("SELECT id, username FROM accounts_user;", config=cfg)
        self.assertTrue(res["success"])
        self.assertEqual(len(res["results"]), 1)
        self.assertEqual(res["results"][0]["username"], "faculty1")
        mock_post.assert_called_once()
        self.assertIn("acc-12345/d1/database/d1-uuid-67890/query", mock_post.call_args[0][0])

    @patch("apps.core.services.d1_service.CloudflareD1Service.execute_query")
    def test_execute_d1_read_mocked(self, mock_exec):
        from apps.core.services.d1_service import execute_d1_read

        mock_exec.return_value = {
            "success": True,
            "results": [{"id": 10, "name": "Computer Science"}],
            "meta": {"duration": 0.4, "rows_read": 1},
        }
        res = execute_d1_read("SELECT id, name FROM accounts_department;", config={"account_id": "a", "database_id": "d", "api_token": "t"})
        self.assertTrue(res["success"])
        self.assertEqual(res["count"], 1)
        self.assertEqual(res["rows"][0]["name"], "Computer Science")

    @patch("apps.core.services.d1_service.CloudflareD1Service.execute_query")
    def test_execute_d1_write_mocked(self, mock_exec):
        from apps.core.services.d1_service import execute_d1_write

        mock_exec.return_value = {
            "success": True,
            "results": [],
            "meta": {"changes": 3, "rows_written": 3, "last_row_id": 99, "duration": 0.8},
        }
        res = execute_d1_write("UPDATE accounts_user SET is_active=1;", config={"account_id": "a", "database_id": "d", "api_token": "t"})
        self.assertTrue(res["success"])
        self.assertEqual(res["changes"], 3)
        self.assertEqual(res["rows_written"], 3)
        self.assertEqual(res["last_row_id"], 99)

    def test_split_sql_script(self):
        from apps.core.services.d1_service import CloudflareD1Service

        script = """
        -- Setup schema
        BEGIN TRANSACTION;
        CREATE TABLE IF NOT EXISTS test_tbl (id INTEGER PRIMARY KEY, title TEXT);
        /* multiline comment */
        CREATE INDEX IF NOT EXISTS idx_test ON test_tbl(title);
        COMMIT;
        """
        stmts = CloudflareD1Service.split_sql_script(script)
        self.assertEqual(len(stmts), 2)
        self.assertIn("CREATE TABLE IF NOT EXISTS test_tbl", stmts[0])
        self.assertIn("CREATE INDEX IF NOT EXISTS idx_test", stmts[1])

    @patch("apps.core.services.d1_service.CloudflareD1Service.execute_query")
    def test_execute_batch_and_script(self, mock_exec):
        from apps.core.services.d1_service import CloudflareD1Service

        mock_exec.return_value = {
            "success": True,
            "all_results": [
                {"meta": {"changes": 1, "rows_written": 1, "rows_read": 0, "duration": 0.2}},
                {"meta": {"changes": 0, "rows_written": 0, "rows_read": 0, "duration": 0.1}},
            ],
            "meta": {"changes": 1, "rows_written": 1, "rows_read": 0, "duration": 0.3},
        }

        svc = CloudflareD1Service(config={"account_id": "a", "database_id": "d", "api_token": "t"})
        res = svc.execute_script("CREATE TABLE a (id INT); CREATE TABLE b (id INT);")
        self.assertTrue(res["success"])
        self.assertEqual(res["total_statements"], 2)

    @patch("apps.core.services.d1_service.CloudflareD1Service.execute_script")
    def test_apply_schema_mocked(self, mock_script):
        from apps.core.services.d1_service import CloudflareD1Service
        import tempfile

        mock_script.return_value = {"success": True, "total_statements": 5, "batches_executed": 1, "changes": 0, "duration": 0.5}
        svc = CloudflareD1Service(config={"account_id": "a", "database_id": "d", "api_token": "t"})

        with tempfile.NamedTemporaryFile("w+", suffix=".sql") as f:
            f.write("CREATE TABLE demo (id INT);")
            f.flush()
            res = svc.apply_schema(schema_file=f.name)
            self.assertTrue(res["success"])
            mock_script.assert_called_once()

    @patch("apps.core.services.d1_service.CloudflareD1Service.execute_read")
    def test_list_tables_and_counts(self, mock_read):
        from apps.core.services.d1_service import CloudflareD1Service

        # First call is list_tables, second is count for t1, third is count for t2
        mock_read.side_effect = [
            {"success": True, "rows": [{"name": "accounts_user"}, {"name": "activities_activity"}]},
            {"success": True, "rows": [{"cnt": 15}]},
            {"success": True, "rows": [{"cnt": 42}]},
        ]

        svc = CloudflareD1Service(config={"account_id": "a", "database_id": "d", "api_token": "t"})
        counts = svc.get_table_counts()
        self.assertEqual(counts["accounts_user"], 15)
        self.assertEqual(counts["activities_activity"], 42)

    @patch("apps.core.services.d1_service.execute_d1_query")
    def test_test_d1_connection_mocked(self, mock_query):
        from apps.core.services.d1_service import test_d1_connection

        mock_query.return_value = {
            "success": True,
            "results": [{"probe": 1, "d1_now": "2026-09-21 04:00:00"}],
        }
        cfg = {"account_id": "acc", "database_id": "db", "api_token": "tok"}
        success, msg = test_d1_connection(cfg)
        self.assertTrue(success)
        self.assertIn("2026-09-21", msg)

    def test_export_django_schema_for_d1(self):
        from apps.core.services.d1_service import export_django_schema_for_d1

        schema_sql = export_django_schema_for_d1()
        self.assertIsInstance(schema_sql, str)
        self.assertIn("Cloudflare D1 SQLite Compatible Database Schema", schema_sql)
        self.assertIn("PRAGMA foreign_keys = ON;", schema_sql)


class D1CliManagementCommandTests(TestCase):
    def test_d1_cli_export_schema_no_wrangler(self):
        import tempfile
        from django.core.management import call_command

        with tempfile.NamedTemporaryFile(suffix=".sql") as tmp:
            out = io.StringIO()
            call_command("d1_cli", f"--export-schema={tmp.name}", stdout=out)
            output = out.getvalue()
            with open(tmp.name, "r") as f:
                content = f.read()
            self.assertIn("Cloudflare D1", content)
            self.assertGreater(len(content), 100)
            self.assertNotIn("wrangler", output.lower())
            self.assertIn("--apply-schema", output)

    @patch("apps.core.management.commands.d1_cli.execute_d1_read")
    def test_d1_cli_read(self, mock_read):
        from django.core.management import call_command

        mock_read.return_value = {
            "success": True,
            "rows": [{"count": 7}],
            "meta": {"duration": 0.2, "rows_read": 1},
        }
        out = io.StringIO()
        call_command("d1_cli", "--read=SELECT count(*) FROM accounts_user;", stdout=out)
        self.assertIn("Query returned 1 row(s)", out.getvalue())

    @patch("apps.core.management.commands.d1_cli.execute_d1_write")
    def test_d1_cli_write(self, mock_write):
        from django.core.management import call_command

        mock_write.return_value = {
            "success": True,
            "changes": 1,
            "rows_written": 1,
            "last_row_id": 5,
            "meta": {"duration": 0.3},
        }
        out = io.StringIO()
        call_command("d1_cli", "--write=UPDATE accounts_user SET is_active=1 WHERE id=5;", stdout=out)
        self.assertIn("Write statement executed successfully", out.getvalue())
        self.assertIn("Changes: 1", out.getvalue())

    @patch("apps.core.management.commands.d1_cli.apply_d1_schema")
    def test_d1_cli_apply_schema(self, mock_apply):
        from django.core.management import call_command

        mock_apply.return_value = {
            "success": True,
            "total_statements": 12,
            "executed_statements": 12,
            "batches_executed": 1,
            "duration": 0.45,
        }
        out = io.StringIO()
        call_command("d1_cli", "--apply-schema", stdout=out)
        self.assertIn("Schema successfully applied to Cloudflare D1", out.getvalue())

    @patch("apps.core.management.commands.d1_cli.CloudflareD1Service")
    def test_d1_cli_tables(self, mock_service_cls):
        from django.core.management import call_command

        mock_svc = MagicMock()
        mock_svc.get_table_counts.return_value = {
            "accounts_user": 10,
            "activities_activity": 25,
        }
        mock_service_cls.return_value = mock_svc

        out = io.StringIO()
        call_command("d1_cli", "--tables", stdout=out)
        output = out.getvalue()
        self.assertIn("accounts_user", output)
        self.assertIn("activities_activity", output)
        self.assertIn("Found 2 table(s)", output)


